#!/usr/bin/env python3
"""Generate the Axiom Rover Mulo master engineering DOCX."""

from __future__ import annotations

import csv
import re
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "docs" / "deliverables"
DOCX_PATH = OUT_DIR / "axiom_mulo_master_engineering_manual.docx"

REVISION = "REV A"
BASELINE_DATE = "2026-06-14"
PROJECT = "Axiom Rover Mulo"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
NAVY = RGBColor(11, 37, 69)
GRAY = RGBColor(85, 85, 85)
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CALL_OUT = "F4F6F9"
SAFETY_FILL = "FFF2CC"
RISK_FILL = "FCE4D6"


ASSETS = {
    "hero": ROOT / "docs" / "diagrams" / "render_mulo_trekking_alpine.png",
    "studio": ROOT / "docs" / "diagrams" / "render_mulo_studio_dark.png",
    "chassis_layout": ROOT / "hardware" / "cad" / "chassis_layout.png",
    "chassis_bracing": ROOT / "hardware" / "cad" / "chassis_bracing.png",
    "wheel": ROOT / "hardware" / "cad" / "wheel_assembly.png",
    "bridge": ROOT / "hardware" / "cad" / "ponte_corazzato_exploded_view.png",
    "drivetrain": ROOT / "hardware" / "cad" / "blueprint_drivetrain_exploded.png",
    "winch": ROOT / "hardware" / "cad" / "blueprint_winch_assembly.png",
    "power_box": ROOT / "hardware" / "cad" / "cad_power_box_assembly.png",
    "rex_power_current": ROOT / "docs" / "enterprise" / "figures" / "rex_expected_power_current.png",
    "rex_rpm_vdc": ROOT / "docs" / "enterprise" / "figures" / "rex_expected_rpm_vdc.png",
    "rex_thermal": ROOT / "docs" / "enterprise" / "figures" / "rex_expected_thermal.png",
    "wireviz": ROOT / "hardware" / "electronics" / "rex_controller" / "generated" / "wireviz" / "rex_controller_harness.png",
    "rex_arch": ROOT / "hardware" / "electronics" / "rex_controller" / "generated" / "graphs" / "rex_power_architecture.png",
    "rex_carrier": ROOT / "hardware" / "mechanical" / "rex_controller_carrier" / "generated" / "rex_controller_carrier_plate.png",
    "battery": ROOT / "docs" / "diagrams" / "plot_battery_discharge_soc.png",
    "motor_curve": ROOT / "docs" / "diagrams" / "plot_motor_torque_speed.png",
    "uwb": ROOT / "docs" / "diagrams" / "plot_uwb_trilateration_error.png",
    "winch_curve": ROOT / "docs" / "diagrams" / "plot_winch_tension_slope.png",
    "cable_thermal": ROOT / "docs" / "research_thesis" / "electrical_cable_thermal_limits.png",
    "nmpc": ROOT / "docs" / "research_thesis" / "rover_nmpc_casadi_performance.png",
}


SOURCE_DIGESTS = [
    ("Specifiche meccaniche e strutturali finali", ROOT / "docs" / "mulo_mk0_hardware_specs.md", 52000),
    ("Decisioni hardware per autonomia trekking", ROOT / "docs" / "mulo_trekking_autonomy_hardware_decisions.md", 36000),
    ("Disegni tecnici e tavole engineering", ROOT / "docs" / "mulo_mk0_engineering_drawings.md", 22000),
    ("Specifiche sintetiche Mulo", ROOT / "docs" / "mulo_specs.md", 14000),
    ("Verricello e stabilita'", ROOT / "docs" / "winch_dynamics.md", 12000),
    ("Ricarica solare e trade-off energia", ROOT / "docs" / "mulo_solar_charging_tradeoff.md", 18000),
    ("Safety case enterprise", ROOT / "docs" / "enterprise" / "safety_case.md", 24000),
    ("V&V master plan", ROOT / "docs" / "enterprise" / "vv_master_plan.md", 22000),
    ("Baseline elettrica 24 V", ROOT / "docs" / "enterprise" / "electrical_baseline_24v.md", 22000),
    ("Range extender readiness", ROOT / "docs" / "enterprise" / "range_extender_readiness.md", 26000),
    ("Manufacturing readiness package", ROOT / "docs" / "enterprise" / "manufacturing_readiness_package.md", 22000),
    ("Data reporting standard", ROOT / "docs" / "enterprise" / "data_reporting_standard.md", 18000),
    ("Digital twin scenarios", ROOT / "docs" / "enterprise" / "digital_twin_scenarios.md", 14000),
    ("Cybersecurity plan", ROOT / "docs" / "enterprise" / "cybersecurity_plan.md", 18000),
    ("BOM meccanica e sottosistemi", ROOT / "hardware" / "bom.md", 22000),
    ("Albero custom 42CrMo4", ROOT / "hardware" / "mechanical" / "extension_shaft_specs.md", 12000),
    ("Istruzioni assemblaggio REX", ROOT / "hardware" / "electronics" / "rex_controller" / "assembly_work_instructions.md", 15000),
]


CSV_TABLES = [
    ("Traceability matrix", ROOT / "docs" / "enterprise" / "traceability_matrix.csv", 40),
    ("Wiring validation matrix", ROOT / "docs" / "enterprise" / "wiring_validation_matrix.csv", 40),
    ("Range extender validation matrix", ROOT / "docs" / "enterprise" / "range_extender_validation_matrix.csv", 40),
    ("BOM candidate shortlist", ROOT / "docs" / "enterprise" / "mulo_bom_candidate_shortlist.csv", 40),
    ("REX controller harness", ROOT / "hardware" / "electronics" / "rex_controller" / "rex_controller_harness.csv", 60),
    ("REX controller netlist", ROOT / "hardware" / "electronics" / "rex_controller" / "rex_controller_netlist.csv", 60),
    ("REX controller BOM", ROOT / "hardware" / "electronics" / "rex_controller" / "rex_controller_bom.csv", 40),
]


def rel(path: Path) -> str:
    try:
        return str(path.relative_to(ROOT)).replace("\\", "/")
    except ValueError:
        return str(path)


def add_field(paragraph, field_code: str) -> None:
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = field_code
    fld_char_separate = OxmlElement("w:fldChar")
    fld_char_separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_begin)
    run._r.append(instr_text)
    run._r.append(fld_char_separate)
    run._r.append(text)
    run._r.append(fld_char_end)


def set_run_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), "120")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Pt(widths_dxa[idx] / 20)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))


def set_cell_text(cell, text, bold=False, color=None, size=9.0):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.08
    r = p.add_run(str(text))
    set_run_font(r, size=size, color=color or RGBColor(0, 0, 0), bold=bold)


def add_table(doc: Document, headers, rows, widths_dxa=None, font_size=8.5, header_fill=LIGHT_BLUE):
    if widths_dxa is None:
        usable = 9360
        widths_dxa = [usable // len(headers)] * len(headers)
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths_dxa)
    for idx, header in enumerate(headers):
        shade_cell(table.rows[0].cells[idx], header_fill)
        set_cell_text(table.rows[0].cells[idx], header, bold=True, color=NAVY, size=font_size)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            set_cell_text(cells[idx], value, size=font_size)
    doc.add_paragraph()
    return table


def add_para(doc: Document, text: str, style=None, bold=False, italic=False, color=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.25
    run = p.add_run(text)
    set_run_font(run, size=10.6, color=color or RGBColor(0, 0, 0), bold=bold, italic=italic)
    return p


def add_bullet(doc: Document, text: str):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.18
    run = p.add_run(text)
    set_run_font(run, size=10.3)
    return p


def add_number(doc: Document, text: str):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.18
    run = p.add_run(text)
    set_run_font(run, size=10.3)
    return p


def add_heading(doc: Document, text: str, level=1):
    p = doc.add_heading(level=level)
    p.paragraph_format.space_before = Pt(18 if level == 1 else 12)
    p.paragraph_format.space_after = Pt(8 if level == 1 else 6)
    for run in p.runs:
        set_run_font(run, size={1: 16, 2: 13, 3: 12}.get(level, 11), color=BLUE if level < 3 else DARK_BLUE, bold=True)
    p.add_run(text)
    for run in p.runs:
        set_run_font(run, size={1: 16, 2: 13, 3: 12}.get(level, 11), color=BLUE if level < 3 else DARK_BLUE, bold=True)
    return p


def add_callout(doc: Document, label: str, text: str, fill=CALL_OUT):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_table_geometry(table, [9360])
    cell = table.rows[0].cells[0]
    shade_cell(cell, fill)
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(f"{label}: ")
    set_run_font(r, size=10.5, color=NAVY, bold=True)
    r = p.add_run(text)
    set_run_font(r, size=10.3, color=RGBColor(0, 0, 0))
    doc.add_paragraph()


def image_width_for(path: Path, max_width=6.15, max_height=6.7) -> float:
    with Image.open(path) as im:
        width_px, height_px = im.size
    ratio = height_px / width_px
    width = min(max_width, max_height / ratio)
    return max(3.0, width)


def add_figure(doc: Document, path: Path, caption: str, source: str | None = None, max_width=6.15, max_height=6.7):
    if not path.exists():
        add_callout(doc, "Figura mancante", f"Asset non trovato: {rel(path)}", fill=RISK_FILL)
        return False
    width = image_width_for(path, max_width=max_width, max_height=max_height)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(8)
    r = cap.add_run(caption)
    set_run_font(r, size=8.8, color=GRAY, italic=True)
    if source:
        src = doc.add_paragraph()
        src.alignment = WD_ALIGN_PARAGRAPH.CENTER
        src.paragraph_format.space_after = Pt(8)
        r = src.add_run(f"Fonte locale: {source}")
        set_run_font(r, size=7.8, color=GRAY)
    return True


def clean_markdown(line: str) -> str:
    line = line.strip()
    line = re.sub(r"!\[([^\]]*)\]\([^)]+\)", r"\1", line)
    line = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", line)
    line = line.replace("**", "").replace("__", "")
    line = line.replace("`", "")
    line = re.sub(r"<[^>]+>", "", line)
    line = re.sub(r"\s+", " ", line)
    return line.strip()


def add_markdown_digest(doc: Document, title: str, path: Path, max_chars: int):
    if not path.exists():
        return
    add_heading(doc, title, 2)
    add_para(doc, f"Sorgente consolidata: {rel(path)}", italic=True, color=GRAY)
    text = path.read_text(encoding="utf-8", errors="ignore")
    consumed = 0
    in_code = False
    last_blank = False
    for raw in text.splitlines():
        if consumed >= max_chars:
            add_callout(doc, "Nota", "Estratto compattato per mantenere il manuale entro il formato di riferimento.", fill=LIGHT_GRAY)
            break
        line = raw.rstrip()
        if line.strip().startswith("```"):
            in_code = not in_code
            continue
        if in_code:
            continue
        if not line.strip():
            last_blank = True
            continue
        if "ESP32" in line or "esp32" in line:
            continue
        if line.strip().startswith("|"):
            continue
        heading_match = re.match(r"^(#{1,4})\s+(.*)$", line)
        if heading_match:
            level = min(3, len(heading_match.group(1)) + 1)
            add_heading(doc, clean_markdown(heading_match.group(2)), level)
            consumed += len(line)
            last_blank = False
            continue
        clean = clean_markdown(line)
        if not clean:
            continue
        if clean.startswith("- "):
            add_bullet(doc, clean[2:])
        elif re.match(r"^\d+[\.\)]\s+", clean):
            add_number(doc, re.sub(r"^\d+[\.\)]\s+", "", clean))
        else:
            # Avoid tiny orphan fragments from markdown formatting.
            if len(clean) < 25 and not last_blank:
                continue
            add_para(doc, clean)
        consumed += len(line)
        last_blank = False


def read_csv_rows(path: Path, max_rows: int):
    if not path.exists():
        return [], []
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        reader = csv.reader(f)
        rows = list(reader)
    if not rows:
        return [], []
    header = rows[0]
    body = rows[1 : 1 + max_rows]
    return header, body


def short(text: str, max_len=120) -> str:
    text = str(text).replace("\n", " ").strip()
    if len(text) <= max_len:
        return text
    return text[: max_len - 1].rstrip() + "..."


def add_csv_table(doc: Document, title: str, path: Path, max_rows: int):
    header, rows = read_csv_rows(path, max_rows)
    if not header:
        return
    add_heading(doc, title, 2)
    add_para(doc, f"Sorgente: {rel(path)}. Righe incluse: {len(rows)}.", italic=True, color=GRAY)
    rows = [[short(cell, 95) for cell in row[: len(header)]] for row in rows]
    if len(header) == 2:
        widths = [2200, 7160]
    elif len(header) == 3:
        widths = [1800, 3300, 4260]
    elif len(header) == 4:
        widths = [1500, 2500, 2800, 2560]
    elif len(header) == 5:
        widths = [1200, 1900, 2500, 2100, 1660]
    else:
        widths = [max(900, 9360 // len(header))] * len(header)
    add_table(doc, [short(h, 45) for h in header], rows, widths_dxa=widths, font_size=7.6)


def add_source_inventory(doc: Document):
    add_heading(doc, "Inventario asset e sorgenti usati", 1)
    files = []
    for pattern in ["*.md", "*.csv", "*.png", "*.svg", "*.pdf", "*.step", "*.stl", "*.dxf", "*.FCStd"]:
        files.extend(ROOT.rglob(pattern))
    selected = []
    include_dirs = [
        ROOT / "docs" / "enterprise",
        ROOT / "docs" / "diagrams",
        ROOT / "docs" / "research_thesis",
        ROOT / "hardware" / "cad",
        ROOT / "hardware" / "electronics" / "rex_controller",
        ROOT / "hardware" / "mechanical",
        ROOT / "hardware" / "firmware",
        ROOT / "reports" / "evidence",
    ]
    for f in sorted(set(files)):
        if any(str(f).startswith(str(d)) for d in include_dirs):
            selected.append(f)
    rows = []
    for f in selected[:120]:
        rows.append([rel(f), f.suffix.lower() or "file", str(f.stat().st_size)])
    add_table(doc, ["Percorso", "Tipo", "Byte"], rows, widths_dxa=[6500, 1300, 1560], font_size=7.4)


def configure_document(doc: Document):
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(10.6)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color in [
        ("Heading 1", 16, BLUE),
        ("Heading 2", 13, BLUE),
        ("Heading 3", 12, DARK_BLUE),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(14 if name == "Heading 1" else 10)
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.keep_with_next = True


def add_running_header_footer(doc: Document):
    section = doc.sections[0]
    header = section.header
    p = header.paragraphs[0]
    p.text = ""
    r = p.add_run(f"{PROJECT} - Master Engineering & Assembly Manual")
    set_run_font(r, size=8.5, color=GRAY, bold=True)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT

    footer = section.footer
    p = footer.paragraphs[0]
    p.text = ""
    r = p.add_run(f"{REVISION} | Baseline {BASELINE_DATE} | Pagina ")
    set_run_font(r, size=8.5, color=GRAY)
    add_field(p, "PAGE")
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def add_cover(doc: Document):
    add_running_header_footer(doc)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(28)
    r = p.add_run("AXIOM ROVER MULO")
    set_run_font(r, size=30, color=NAVY, bold=True)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p = doc.add_paragraph()
    r = p.add_run("Master Engineering & Assembly Manual")
    set_run_font(r, size=17, color=DARK_BLUE, bold=True)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p = doc.add_paragraph()
    r = p.add_run("Fotografia tecnica del prototipo finale da assemblare")
    set_run_font(r, size=12.5, color=GRAY, italic=True)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    add_figure(
        doc,
        ASSETS["hero"],
        "Vista di riferimento del rover Mulo in scenario trekking.",
        source=rel(ASSETS["hero"]),
        max_width=6.4,
        max_height=3.5,
    )

    rows = [
        ["Documento", "Manuale ingegneristico, assemblaggio e validazione"],
        ["Progetto", PROJECT],
        ["Revisione", REVISION],
        ["Baseline", BASELINE_DATE],
        ["Ambito", "Meccanica, elettrico, software, firmware, safety, test, BOM, readiness"],
        ["Nota", "Non autorizza ordini custom o test outdoor senza chiusura dei gate P0."],
    ]
    add_table(doc, ["Campo", "Valore"], rows, widths_dxa=[1800, 7560], font_size=9.0, header_fill=LIGHT_GRAY)
    add_callout(
        doc,
        "Blocco P0",
        "Il rover non deve essere assemblato o provato in outdoor finche' cablaggi potenza, freni, safety fisica, dump load, logging e test banco non hanno evidenza firmata.",
        fill=SAFETY_FILL,
    )
    doc.add_page_break()


def add_static_toc(doc: Document):
    add_heading(doc, "Indice operativo", 1)
    rows = [
        ["1", "Executive overview", "Missione, baseline e decisioni finali"],
        ["2", "System architecture", "Domini e flussi principali"],
        ["3", "Requisiti e tracciabilita'", "MULO-SYS, V&V, evidence"],
        ["4", "Meccanica completa", "Telaio, ruote, alberi, snodo, freni"],
        ["5", "Drivetrain e trazione", "Motori, VESC, 4WD, modalita' degradata"],
        ["6", "Verricello", "Trekking helper, riduzioni, anti-ribaltamento"],
        ["7", "Elettrico e cablaggi", "24 V, protezioni, WireViz, test isolamento"],
        ["8", "Energia e REX", "LiFePO4, GX50, STM32 REX, grafici banco"],
        ["9", "Safety", "STM32 safety MCU, E-stop, CRC, fail-closed"],
        ["10", "Software e firmware", "ROS 2, profili launch, PlatformIO"],
        ["11", "Autonomia e sensori", "UWB, ToF, IMU, LiDAR/camera, digital twin"],
        ["12", "Assemblaggio e V&V", "Sequenze, checklist, acceptance"],
        ["13", "BOM e appendici", "Procurement, toolchain, sorgenti, glossario"],
    ]
    add_table(doc, ["#", "Sezione", "Contenuto"], rows, widths_dxa=[650, 3100, 5610], font_size=8.6)
    add_callout(
        doc,
        "Uso del documento",
        "Le parti iniziali servono a stakeholder e revisori; le parti centrali sono manuale tecnico; le appendici mantengono evidenza, sorgenti e tabelle per officina e V&V.",
        fill=CALL_OUT,
    )
    doc.add_page_break()


def add_subsystem_card(doc: Document, title: str, purpose: str, configuration: str, interfaces: str, risks: str, validation: str):
    add_heading(doc, title, 2)
    add_table(
        doc,
        ["Voce", "Descrizione"],
        [
            ["Funzione", purpose],
            ["Configurazione finale", configuration],
            ["Interfacce", interfaces],
            ["Rischi tecnici", risks],
            ["Validazione richiesta", validation],
        ],
        widths_dxa=[2100, 7260],
        font_size=8.6,
        header_fill=LIGHT_BLUE,
    )


def add_main_body(doc: Document):
    add_heading(doc, "1. Executive overview", 1)
    add_para(
        doc,
        "Mulo e' un rover da trekking e supporto logistico costruito attorno a una piattaforma 4WD robusta, con telaio custom, drivetrain rinforzato, verricello assistito, batteria 24 V LiFePO4, safety separata e software ROS 2 modulare.",
    )
    add_para(
        doc,
        "La fotografia finale trattata in questo manuale non e' il percorso storico del progetto: e' la baseline tecnica da portare verso assemblaggio, validazione banco e readiness industriale.",
    )
    add_callout(
        doc,
        "Decisione baseline",
        "Safety e controllo critico finale sono su STM32 Nucleo e ROS 2. I riferimenti ESP32 presenti in documenti storici restano legacy o ausiliari non safety-critical.",
        fill=LIGHT_BLUE,
    )
    add_figure(doc, ASSETS["studio"], "Render tecnico/studio del rover Mulo.", source=rel(ASSETS["studio"]))
    add_table(
        doc,
        ["Dominio", "Baseline", "Stato"],
        [
            ["Meccanica", "Telaio custom, ruote, alberi, UCF204, ponte corazzato, verricello", "Da validare con tavole quotate e interferenze"],
            ["Elettrico", "24 V nominale, rami protetti, dump load, harness WireViz", "Baseline pronta, rating finali da chiudere"],
            ["Energia", "LiFePO4 24 V + REX Honda GX50 opzionale", "REX disabilitato finche' banco non firmato"],
            ["Safety", "STM32 safety MCU, E-stop, CRC, heartbeat, consenso motori", "Fail-closed impostato, HIL fisico da completare"],
            ["Software", "ROS 2 package modulari e profili sim/bench/field/production", "Test logici verdi, colcon richiede VS Developer Prompt"],
            ["V&V", "Requisiti, traceability, evidence pack, checklist", "Processo impostato, prove fisiche aperte"],
        ],
        widths_dxa=[1700, 4600, 3060],
    )

    add_heading(doc, "2. System architecture", 1)
    add_para(
        doc,
        "L'architettura e' organizzata per domini: struttura e locomozione, potenza e ricarica, safety real-time, autonomia, logging e validazione. La separazione tra computer principale e MCU safety e' intenzionale: nessun comportamento autonomo deve poter superare il consenso fisico dei motori.",
    )
    add_figure(doc, ASSETS["chassis_layout"], "Layout telaio e ingombri principali.", source=rel(ASSETS["chassis_layout"]))
    add_subsystem_card(
        doc,
        "Dominio meccanico",
        "Sostiene carico, trazione, protezioni, manutenzione e trasporto.",
        "Telaio custom con ruote motrici, assi supportati, box elettronica e predisposizione per verricello/REX.",
        "Interfaccia verso motori, freni, supporti sensori, box IP, cablaggi e punti di manutenzione.",
        "Interferenze, giochi non controllati, vibrazioni, accesso manutenzione insufficiente, massa fuori target.",
        "Review CAD, tavole quotate, controllo interferenze, prove statiche e vibrazione manuale.",
    )
    add_subsystem_card(
        doc,
        "Dominio elettrico",
        "Distribuisce energia 24 V e garantisce separazione fra potenza, logica, safety e ricarica.",
        "Batteria LiFePO4, fusibile principale, sezionatore, rami protetti, DC/DC, dump load, CAN e connettori IP.",
        "BMS, VESC, STM32 safety, STM32 REX, Jetson/ROS, sensori e carichi ausiliari.",
        "Backfeed, sovracorrente, masse rumorose, connettori non IP, cavi non dimensionati.",
        "Continuita', isolamento, test ramo per ramo, misura caduta tensione e fault injection.",
    )
    add_subsystem_card(
        doc,
        "Dominio software",
        "Coordina controllo, energia, autonomia, safety monitoring, logging e report.",
        "ROS 2 modulare con package dedicati e profili sim/bench/field/production.",
        "Topic SafetyCommand, McuStatus, RangeExtenderStatus, VESC/CAN, sensor fusion e logging.",
        "Simulazione attiva in profilo reale, dipendenze non dichiarate, mismatch interfacce.",
        "pytest, compileall, colcon in Developer Prompt, replay dati, test perdita MCU/VESC.",
    )

    add_heading(doc, "3. Requisiti, traceability e readiness", 1)
    add_para(
        doc,
        "Il progetto e' trattato come sistema verificabile. Ogni requisito deve avere design, test, evidenza, BOM e rischio associato. La matrice di tracciabilita' e' il contratto fra progetto e assemblaggio.",
    )
    add_callout(
        doc,
        "Gate acquisti",
        "Nessuna parte costosa o lavorazione custom deve essere ordinata se il requisito collegato ha evidenza P0 mancante, rating TBD o schema non validato.",
        fill=SAFETY_FILL,
    )
    add_csv_table(doc, "Matrice requisiti -> design -> test -> evidenza", ROOT / "docs" / "enterprise" / "traceability_matrix.csv", 28)

    add_heading(doc, "4. Meccanica completa", 1)
    add_para(
        doc,
        "La meccanica e' il collo di bottiglia prima dell'assemblaggio fisico: il pacchetto richiede tavole quotate, assiemi, quote funzionali, materiali, tolleranze, accesso manutenzione e controllo interferenze.",
    )
    add_figure(doc, ASSETS["chassis_bracing"], "Struttura telaio e bracing.", source=rel(ASSETS["chassis_bracing"]))
    add_figure(doc, ASSETS["wheel"], "Assieme ruota e supporti.", source=rel(ASSETS["wheel"]))
    add_figure(doc, ASSETS["bridge"], "Ponte corazzato e drivetrain rinforzato.", source=rel(ASSETS["bridge"]))
    add_subsystem_card(
        doc,
        "Telaio e geometrie",
        "Fornire base rigida e ispezionabile per ruote, power box, sensori e kit manutenzione.",
        "Telaio custom con rinforzi, protezioni, predisposizione quick-split/trasporto e fissaggi box.",
        "Fori/staffe, passaggi cavo, masse, supporti motore, freni, snodo e piano superiore.",
        "Disallineamento, deformazioni, interferenze con ruote/cinghie/cavi, manutenzione difficoltosa.",
        "Tavola quotata, assieme CAD completo, controllo interferenze e revisione officina.",
    )
    add_subsystem_card(
        doc,
        "Ruote, alberi e supporti",
        "Trasmettere coppia senza caricare direttamente i motori con flessione e urti.",
        "Alberi custom, supporti UCF204, flange/mozzi, disco/freno e ruote rinforzate.",
        "Motore, albero, cuscinetti, disco freno, mozzo, ruota e sensori velocita'.",
        "Runout, gioco cuscinetti, concentratori di tensione, montaggio non coassiale.",
        "Misura runout, coppie serraggio, ispezione filetti, prova ruote sollevate.",
    )

    add_heading(doc, "5. Drivetrain, trazione e frenata", 1)
    add_figure(doc, ASSETS["drivetrain"], "Vista esplosa drivetrain.", source=rel(ASSETS["drivetrain"]))
    add_figure(doc, ASSETS["motor_curve"], "Curva motore/coppia usata per dimensionamento preliminare.", source=rel(ASSETS["motor_curve"]))
    add_para(
        doc,
        "La trazione 4WD deve restare controllabile anche in degrado. Il software deve limitare velocita' e corrente se una ruota o un controller non risponde; la meccanica deve permettere limp-home manuale senza smontaggi distruttivi.",
    )
    add_callout(
        doc,
        "Safety drivetrain",
        "La frenata di sicurezza non deve dipendere dal solo software applicativo: E-stop fisico, consenso motori e freno meccanico restano indipendenti dalla logica di autonomia.",
        fill=SAFETY_FILL,
    )

    add_heading(doc, "6. Verricello e trekking helper", 1)
    add_figure(doc, ASSETS["winch"], "Assieme verricello custom.", source=rel(ASSETS["winch"]))
    add_figure(doc, ASSETS["winch_curve"], "Tensione verricello e pendenza: envelope di rischio back-flip.", source=rel(ASSETS["winch_curve"]))
    add_subsystem_card(
        doc,
        "Verricello assistito",
        "Aiutare l'utente in salita e gestire traino/guinzaglio con controllo di ammettenza.",
        "Riduzione meccanica, freno/pin-lock, cella carico, cavo e gestione tensione.",
        "Nodo ROS winch, sensori tensione, safety watchdog, comandi trazione e stato pendenza.",
        "Back-flip, sovratensione cavo, rilascio improvviso, freno non fail-safe.",
        "Banco statico, test tensione nota, test freno, test pendenza simulata e stop rapido.",
    )

    add_heading(doc, "7. Elettrico, cablaggi e harness", 1)
    add_figure(doc, ASSETS["power_box"], "Power box e disposizione elettronica.", source=rel(ASSETS["power_box"]))
    add_figure(doc, ASSETS["wireviz"], "Harness REX WireViz: connettori, colori, segnali e note.", source=rel(ASSETS["wireviz"]), max_width=5.0, max_height=7.0)
    add_figure(doc, ASSETS["cable_thermal"], "Limiti termici e derating cavi.", source=rel(ASSETS["cable_thermal"]))
    add_para(
        doc,
        "Il cablaggio finale deve essere una distinta costruibile, non solo uno schema concettuale. Ogni filo richiede sezione, colore, connettore, pin, rating, protezione, strain relief, test continuita' e test isolamento.",
    )
    add_csv_table(doc, "Wiring validation matrix", ROOT / "docs" / "enterprise" / "wiring_validation_matrix.csv", 34)

    add_heading(doc, "8. Alimentazione e range extender Honda GX50", 1)
    add_para(
        doc,
        "Il range extender e' opzionale e separato dalla safety primaria. Non abilita test outdoor finche' rover, freni, cablaggi, dump load, logging e procedure fuel/CO/fire non sono validati.",
    )
    add_figure(doc, ASSETS["rex_arch"], "Architettura potenza REX: GX50 -> BLDC -> DC link -> buck -> batteria.", source=rel(ASSETS["rex_arch"]))
    add_figure(doc, ASSETS["rex_power_current"], "Envelope REX potenza/corrente per prova banco.", source=rel(ASSETS["rex_power_current"]))
    add_figure(doc, ASSETS["rex_rpm_vdc"], "Envelope REX RPM e DC link.", source=rel(ASSETS["rex_rpm_vdc"]))
    add_figure(doc, ASSETS["rex_thermal"], "Envelope termico REX preliminare.", source=rel(ASSETS["rex_thermal"]))
    add_figure(doc, ASSETS["rex_carrier"], "Piastra carrier STM32 REX generata con OpenSCAD/FreeCAD.", source=rel(ASSETS["rex_carrier"]))
    add_csv_table(doc, "Range extender validation matrix", ROOT / "docs" / "enterprise" / "range_extender_validation_matrix.csv", 30)

    add_heading(doc, "9. Safety architecture", 1)
    add_para(
        doc,
        "La safety e' progettata fail-closed: in assenza di comando valido, heartbeat, CRC, seriale/CAN o consenso fisico, il sistema deve portarsi a comando motori nullo e stato sicuro.",
    )
    add_table(
        doc,
        ["Elemento", "Regola finale"],
        [
            ["SafetyCommand", "seq, stamp_ms, mode, velocita', timeout, E-stop, enable motors, max current, crc16"],
            ["McuStatus", "seq_ack, heartbeat_age_ms, safety_state, motor_consent, fault_code, sensor_validity, crc16"],
            ["Heartbeat", "timeout massimo 500 ms prima di chiusura safety"],
            ["E-stop", "fisico, prioritario, indipendente dalla logica AI/follow-me"],
            ["Range extender", "fault critico o kill request non possono essere ignorati da ROS"],
        ],
        widths_dxa=[2200, 7160],
    )
    add_callout(
        doc,
        "Principio",
        "L'autonomia puo' proporre comandi, ma non possiede l'autorita' finale sul movimento. L'autorita' finale passa da safety MCU, E-stop e consenso motori.",
        fill=SAFETY_FILL,
    )

    add_heading(doc, "10. Software ROS 2 e firmware STM32", 1)
    add_para(
        doc,
        "Il workspace ROS 2 contiene package separati per controllo, navigazione, power, winch, safety, system bridge e interfacce. Il bringup finale deve distinguere nettamente simulazione, banco, campo e produzione.",
    )
    add_table(
        doc,
        ["Package/area", "Ruolo"],
        [
            ["rover_control", "VESC driver, kinematics, controllo basso livello"],
            ["rover_interfaces", "SafetyCommand, McuStatus, RangeExtenderStatus e messaggi custom"],
            ["rover_power", "Battery EKF, ECMS, supervisor energia e REX"],
            ["rover_safety", "watchdog, E-stop, stabilita' e fault aggregation"],
            ["rover_system", "hardware bridge, CRC, stato MCU, sensor fusion base"],
            ["rover_navigation", "follow-me, shared autonomy, target detection"],
            ["rover_winch", "gestione verricello e tensione"],
            ["firmware/stm32_safety_mcu", "consenso motori, CRC, heartbeat, fail-closed"],
            ["firmware/stm32_rex_controller", "kill ignition, throttle, RPM/Vdc/Icharge/temp, aux_crc16"],
        ],
        widths_dxa=[3100, 6260],
    )
    add_callout(
        doc,
        "Verifica locale",
        "pytest ha dato 21 passed; PlatformIO ha compilato safety e REX su F446RE/F103RB. colcon su PowerShell standard richiede Visual Studio Developer Prompt per VisualStudioVersion.",
        fill=LIGHT_BLUE,
    )

    add_heading(doc, "11. Autonomia, sensori e digital twin", 1)
    add_figure(doc, ASSETS["uwb"], "Errore UWB/trilaterazione: envelope di analisi.", source=rel(ASSETS["uwb"]))
    add_figure(doc, ASSETS["nmpc"], "Risultati preliminari NMPC/CasADi.", source=rel(ASSETS["nmpc"]))
    add_para(
        doc,
        "Follow-me, UWB, ToF, IMU, camera/LiDAR e digital twin devono restare fuori dalla safety primaria finche' logging, fault injection e test ripetibili non dimostrano robustezza. Il digital twin deve servire soprattutto a ripetere scenari di perdita heartbeat, stallo, pendenza, UWB loss e fault VESC.",
    )

    add_heading(doc, "12. Assemblaggio, V&V ed evidence", 1)
    add_para(
        doc,
        "L'assemblaggio finale procede per energia crescente: prima logica e continuita', poi banco a ruote sollevate, poi sottosistemi energetici limitati, infine test field controllati. Ogni passaggio produce evidenza in reports/evidence.",
    )
    for step in [
        "Generare artefatti: KiCad/WireViz/Graphviz/OpenSCAD/FreeCAD e allegare versione toolchain.",
        "Fare review meccanica: quote, tolleranze, interferenze, accesso manutenzione e passaggi cavo.",
        "Cablaggio a bassa energia: continuita', isolamento, pinout, strain relief e colori.",
        "Banco ruote sollevate: consenso motori default off, E-stop, perdita seriale/CAN/VESC, sovracorrente.",
        "Banco REX spento: kill ignition, servo idle, RPM/Vdc/Icharge simulati, aux_crc16.",
        "Banco REX acceso solo con guardie, CO monitor, estintore, fuel shutoff e carico resistivo.",
        "Field test solo con safety fisica, freni, cablaggio potenza, dump load e logging firmati.",
    ]:
        add_number(doc, step)
    add_csv_table(doc, "BOM candidate shortlist", ROOT / "docs" / "enterprise" / "mulo_bom_candidate_shortlist.csv", 28)

    add_heading(doc, "13. Toolchain e rigenerazione", 1)
    add_para(
        doc,
        "Il pacchetto e' pensato per essere rigenerabile. Gli artefatti di elettronica e meccanica non sono screenshot isolati: hanno sorgenti, comandi e output verificabili.",
    )
    add_csv_table(doc, "REX controller harness", ROOT / "hardware" / "electronics" / "rex_controller" / "rex_controller_harness.csv", 60)


def add_appendices(doc: Document):
    doc.add_page_break()
    add_heading(doc, "Appendici tecniche consolidate", 1)
    add_callout(
        doc,
        "Criterio appendici",
        "Le appendici non raccontano la storia del progetto: raccolgono materiale tecnico ancora utile per assemblaggio, validazione e revisione finale.",
        fill=CALL_OUT,
    )
    for title, path, max_chars in SOURCE_DIGESTS:
        add_markdown_digest(doc, title, path, max_chars)
        doc.add_page_break()

    add_heading(doc, "Appendici tabellari", 1)
    for title, path, max_rows in CSV_TABLES:
        add_csv_table(doc, title, path, max_rows)
    add_source_inventory(doc)

    add_heading(doc, "Glossario operativo", 1)
    glossary = [
        ["P0", "Blocco prima di assemblaggio, ordine o test ad alta energia."],
        ["Fail-closed", "Comportamento sicuro in assenza di comando, heartbeat o sensore valido."],
        ["REX", "Range extender Honda GX50 con controller STM32 dedicato."],
        ["Evidence pack", "Cartella con manifest, log, report, figure, foto e review del test."],
        ["WireViz", "Sorgente e output per cablaggio costruibile con connettori, pin e colori."],
        ["ERC/DRC", "Electrical/design rule check KiCad per schema e PCB."],
        ["HIL", "Hardware-in-the-loop: test firmware con I/O reali o simulati elettricamente."],
        ["SIL", "Software-in-the-loop: test logici e replay senza hardware fisico."],
        ["DSM/ZMP", "Metriche di stabilita' dinamica e prevenzione ribaltamento."],
        ["ECMS", "Strategia di gestione energia ibrida per ripartire potenza batteria/generatore."],
    ]
    add_table(doc, ["Termine", "Definizione"], glossary, widths_dxa=[1900, 7460])


def save_content_map():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    path = OUT_DIR / "axiom_mulo_master_content_map.csv"
    rows = [["type", "title", "path"]]
    rows.extend(["image", key, rel(value)] for key, value in ASSETS.items() if value.exists())
    rows.extend(["source_digest", title, rel(path)] for title, path, _ in SOURCE_DIGESTS if path.exists())
    rows.extend(["csv_table", title, rel(path)] for title, path, _ in CSV_TABLES if path.exists())
    with path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.writer(f)
        writer.writerows(rows)


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    save_content_map()
    doc = Document()
    configure_document(doc)
    doc.core_properties.title = "Axiom Rover Mulo - Master Engineering & Assembly Manual"
    doc.core_properties.subject = "Fotografia tecnica del prototipo finale da assemblare"
    doc.core_properties.author = "Codex / Axiom Rover Workspace"
    doc.core_properties.keywords = "Axiom Rover, Mulo, ROS 2, STM32, safety, assembly, engineering"
    doc.core_properties.comments = f"{REVISION}, baseline {BASELINE_DATE}"

    add_cover(doc)
    add_static_toc(doc)
    add_main_body(doc)
    add_appendices(doc)
    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    main()
